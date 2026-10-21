"""数据初始化审计（任务 5.5）。

检查项：
1. 样例源文件是否带 ``source: 示例数据（待替换为学校真实文件）`` 标记；
2. 数据库文档是否已入库、已发布、切片数一致；
3. 引用是否能回溯到切片与文档（孤立引用为 0）；
4. 源文件与切片中是否残留身份证号 / 手机号 / 学号 / 工号。

用法::

    python tools/verify_data_init.py
    python tools/verify_data_init.py --kb-dir ../knowledge-base --out artifacts/acceptance/data-init-report.md
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

SERVICE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SERVICE_DIR))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # noqa: BLE001
    pass

from sqlalchemy import func, select  # noqa: E402

from app.config import settings  # noqa: E402
from app.db import kb_repo  # noqa: E402
from app.db import models as M  # noqa: E402
from app.db.session import db_session  # noqa: E402

MARKER = "示例数据（待替换为学校真实文件）"
DEFAULT_KB_DIR = SERVICE_DIR.parent / "knowledge-base"
DEFAULT_OUT = SERVICE_DIR / "artifacts" / "acceptance" / "data-init-report.md"

SOURCE_RE = re.compile(r"^source:\s*(.+?)\s*$", re.MULTILINE)
PII_PATTERNS: dict[str, re.Pattern[str]] = {
    "身份证号": re.compile(
        r"(?<!\d)[1-9]\d{5}(?:19|20)\d{2}(?:0[1-9]|1[0-2])"
        r"(?:0[1-9]|[12]\d|3[01])\d{3}[\dXx](?!\d)"
    ),
    "手机号": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "学号": re.compile(r"(?<!\d)20\d{2}\d{6,8}(?!\d)"),
    "工号": re.compile(r"(?<![A-Za-z0-9])T\d{5}(?![0-9])"),
}
SOURCE_EXTS = {".md", ".txt", ".csv", ".json", ".xlsx"}
SKIP_SOURCE_NAMES = {"README.md", "收集清单.md"}
SKIP_DIR_NAMES = {"00-知识标准"}


def _mask(value: str) -> str:
    if len(value) <= 5:
        return value[:1] + "*" * (len(value) - 1)
    return value[:3] + "*" * (len(value) - 5) + value[-2:]


def _read_marker(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext in {".md", ".txt"}:
            text = path.read_text(encoding="utf-8", errors="replace")[:3000]
            match = SOURCE_RE.search(text)
            return match.group(1).strip() if match else ""
        if ext == ".json":
            obj = json.loads(path.read_text(encoding="utf-8"))
            items = obj if isinstance(obj, list) else [obj]
            for item in items:
                if isinstance(item, dict) and str(item.get("source", "")).strip():
                    return str(item["source"]).strip()
            return ""
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                return (wb.properties.description or "").strip()
            finally:
                wb.close()
    except Exception as exc:  # noqa: BLE001
        return f"<读取失败：{exc}>"
    return ""


def _read_source_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            cells: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells.extend(str(c) for c in row if c is not None)
            return "\n".join(cells)
        finally:
            wb.close()
    return path.read_text(encoding="utf-8", errors="replace")


def _source_files(kb_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(kb_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SOURCE_EXTS:
            continue
        if path.name in SKIP_SOURCE_NAMES or any(part in SKIP_DIR_NAMES for part in path.parts):
            continue
        files.append(path)
    return files


def _scan_pii(label: str, text: str) -> list[dict[str, str]]:
    hits: list[dict[str, str]] = []
    for kind, pattern in PII_PATTERNS.items():
        for match in pattern.finditer(text):
            hits.append({"scope": label, "type": kind, "sample": _mask(match.group(0))})
    return hits


def _db_audit() -> dict[str, Any]:
    docs = kb_repo.list_docs()
    doc_rows: list[dict[str, Any]] = []
    published_ids: set[int] = set()
    pii_hits: list[dict[str, str]] = []
    for doc in docs:
        chunks = kb_repo.list_chunks(doc["docId"])
        if doc["status"] == 2:
            published_ids.add(doc["docId"])
        for chunk in chunks:
            pii_hits.extend(_scan_pii(f"{doc['title']} / {chunk['chunkId']}", chunk["text"]))
        doc_rows.append(
            {
                "docId": doc["docId"],
                "title": doc["title"],
                "status": doc["statusName"],
                "format": doc["fileFormat"],
                "fileKey": str(doc.get("fileKey") or "").replace("\\", "/"),
                "chunkCount": doc["chunkCount"],
                "actualChunks": len(chunks),
                "roles": ",".join(doc["visibleRoles"]) or "-",
            }
        )

    with db_session() as conn:
        citation_total = conn.scalar(select(func.count()).select_from(M.AnswerCitation)) or 0
        citation_orphans = (
            conn.scalar(
                select(func.count())
                .select_from(M.AnswerCitation)
                .outerjoin(M.KbChunk, M.AnswerCitation.chunk_id == M.KbChunk.chunk_id)
                .where(M.KbChunk.chunk_id.is_(None))
            )
            or 0
        )
        cited_docs = (
            conn.scalar(select(func.count(func.distinct(M.AnswerCitation.doc_id)))) or 0
        )
        cited_unpublished = (
            conn.scalar(
                select(func.count(func.distinct(M.AnswerCitation.doc_id))).where(
                    M.AnswerCitation.doc_id.notin_(published_ids)
                )
            )
            if published_ids
            else 0
        )

    return {
        "docs": doc_rows,
        "published": len(published_ids),
        "total": len(doc_rows),
        "chunkTotal": sum(r["actualChunks"] for r in doc_rows),
        "citationTotal": int(citation_total),
        "citationOrphans": int(citation_orphans),
        "citedDocs": int(cited_docs),
        "citedUnpublished": int(cited_unpublished),
        "pii": pii_hits,
    }


def _build_report(kb_dir: Path, source_rows: list[dict[str, Any]], db: dict[str, Any]) -> tuple[str, int, int, int]:
    fails = 0
    warns = 0
    passes = 0

    for row in source_rows:
        if row["marker"] == MARKER:
            row["result"] = "PASS"
            passes += 1
        elif row["marker"].startswith("<读取失败"):
            row["result"] = "WARN"
            warns += 1
        else:
            row["result"] = "FAIL"
            fails += 1

    if db["docs"]:
        for row in db["docs"]:
            if row["chunkCount"] != row["actualChunks"] or row["actualChunks"] == 0:
                row["result"] = "FAIL"
                fails += 1
            else:
                row["result"] = "PASS"
                passes += 1

        if db["citationOrphans"]:
            fails += 1
        else:
            passes += 1

        if db["citedUnpublished"]:
            warns += 1

    if db["pii"]:
        fails += 1
    else:
        passes += 1

    not_ingested: list[str] = []
    no_source: list[str] = []
    if db["docs"]:
        source_files = {r["file"] for r in source_rows}
        db_files = {
            r["fileKey"] or str(r["title"])
            for r in db["docs"]
        }
        not_ingested = sorted(source_files - db_files)
        no_source = sorted(db_files - source_files)
        warns += len(not_ingested) + len(no_source)

    verdict = "FAIL" if fails else ("WARN" if warns else "PASS")
    lines = [
        "# 数据初始化审计报告（任务 5.5）",
        "",
        f"> 生成时间：{datetime.now():%Y-%m-%d %H:%M:%S}｜知识目录：`{kb_dir}`｜存储：`{settings.store_label}`",
        f"> 结果：**{verdict}**（PASS {passes} / WARN {warns} / FAIL {fails}）",
        "",
        "## 结论",
        "",
        f"- 数据库文档 {db['total']} 篇，已发布 {db['published']} 篇，切片 {db['chunkTotal']} 条；"
        f"引用 {db['citationTotal']} 条，孤立引用 {db['citationOrphans']} 条，覆盖文档 {db['citedDocs']} 篇。",
    ]
    if db["total"] == 0 and not db["docs"]:
        lines.append("- 本次为源文件预检，未连接数据库；入库后需再跑一次全量审计。")
    lines.extend(
        [
            f"- 样例标记：{sum(1 for r in source_rows if r['marker'] == MARKER)}/{len(source_rows)} 个源文件已标记。",
            f"- PII 扫描：源文件与切片共命中 **{len(db['pii'])}** 处（身份证/手机号/学号/工号）。",
        ]
    )
    if not_ingested:
        lines.append(f"- 未入库源文件：{'、'.join(not_ingested)}。")
    if no_source:
        lines.append(f"- 无对应源文件的库内文档：{'、'.join(no_source)}。")
    lines.extend(
        [
            "",
            "## 1. 源文件样例标记",
            "",
            "| 文件 | 标记 | 结果 |",
            "|------|------|------|",
        ]
    )
    for row in source_rows:
        marker = row["marker"] or "（缺失）"
        lines.append(f"| `{row['file']}` | {marker} | {row['result']} |")

    lines.extend(
        [
            "",
            "## 2. 数据库文档与切片",
            "",
            "| docId | 标题 | 状态 | 格式 | 切片（声明/实际） | 可见角色 | 结果 |",
            "|------:|------|------|------|------------------|----------|------|",
        ]
    )
    for row in db["docs"]:
        lines.append(
            f"| {row['docId']} | {row['title']} | {row['status']} | {row['format']} | "
            f"{row['chunkCount']}/{row['actualChunks']} | {row['roles']} | {row['result']} |"
        )

    lines.extend(
        [
            "",
            "## 3. 引用可追溯",
            "",
            f"- 引用总数：{db['citationTotal']}；孤立引用（切片不存在）：{db['citationOrphans']}。",
            f"- 被引用文档：{db['citedDocs']} 篇；引用到未发布文档：{db['citedUnpublished']} 篇。",
            "",
            "## 4. PII 扫描",
            "",
        ]
    )
    if db["pii"]:
        lines.extend(["| 范围 | 类型 | 脱敏样例 |", "|------|------|----------|"])
        for hit in db["pii"][:50]:
            lines.append(f"| {hit['scope']} | {hit['type']} | `{hit['sample']}` |")
        if len(db["pii"]) > 50:
            lines.append(f"| … | 其余 {len(db['pii']) - 50} 处 | — |")
    else:
        lines.append("未发现身份证号、手机号、学号或工号残留。")

    lines.extend(
        [
            "",
            "## 5. 待办与口径",
            "",
            "- 学校真实文档未到位：本报告只覆盖样例数据；替换流程见《数据初始化与真实文档替换方案》。",
            "- 未入库源文件属于知识运营待办，不影响已发布问答；上线前需确认是否纳入本期知识范围。",
            "- PII 扫描为正则初筛，真实文档入库前仍需人工复核；二进制格式（docx/pdf）请在脱敏阶段单独检查。",
            "",
        ]
    )
    return "\n".join(lines), passes, warns, fails


def main() -> int:
    parser = argparse.ArgumentParser(description="数据初始化审计（5.5）")
    parser.add_argument("--kb-dir", default=str(DEFAULT_KB_DIR), help="知识库源文件目录")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="报告输出路径")
    parser.add_argument(
        "--sources-only",
        action="store_true",
        help="只审计源文件（用于真实文档入库前的脱敏预检，不连接数据库）",
    )
    args = parser.parse_args()

    kb_dir = Path(args.kb_dir).resolve()
    if not kb_dir.is_dir():
        print(f"知识目录不存在：{kb_dir}", file=sys.stderr)
        return 2
    if not args.sources_only and settings.store_label != "mysql":
        print(f"当前存储为 {settings.store_label}，5.5 审计仅支持 MySQL 生产库。", file=sys.stderr)
        return 2

    source_rows = [
        {
            "file": str(path.relative_to(kb_dir)).replace("\\", "/"),
            "marker": _read_marker(path),
        }
        for path in _source_files(kb_dir)
    ]
    source_pii: list[dict[str, str]] = []
    for path in _source_files(kb_dir):
        source_pii.extend(
            _scan_pii(str(path.relative_to(kb_dir)).replace("\\", "/"), _read_source_text(path))
        )

    if args.sources_only:
        db: dict[str, Any] = {
            "docs": [],
            "published": 0,
            "total": 0,
            "chunkTotal": 0,
            "citationTotal": 0,
            "citationOrphans": 0,
            "citedDocs": 0,
            "citedUnpublished": 0,
            "pii": source_pii,
        }
    else:
        db = _db_audit()
        db["pii"] = source_pii + db["pii"]

    report, passes, warns, fails = _build_report(kb_dir, source_rows, db)
    if args.sources_only:
        report = report.replace(
            "> 结果：", "> 模式：仅源文件预检（未连接数据库）\n> 结果：", 1
        )
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(report, encoding="utf-8")

    print(f"数据初始化审计：PASS {passes} / WARN {warns} / FAIL {fails}")
    print(f"报告：{out_path}")
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
