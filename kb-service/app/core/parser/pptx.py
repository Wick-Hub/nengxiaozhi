"""PPTX 解析（任务 1.10）。

原型 ``kb-engine/src/parse/index.js`` 的 ``SUPPORTED`` 只有 8 种格式，**没有** pptx，
本模块是**新增能力**，无 Node 侧可回归。

结构约定（与 xlsx 的「工作表：X」风格一致，便于人工核对与预览对比）：

* 每页 → ``heading(level=1)`` ``幻灯片 N``；
* 页标题占位符 → ``heading(level=2)``；
* 文本框 → ``para``（同一文本框内的段落用换行连接，缩进层级用两个空格表示）；
* 表格 → ``table``（Markdown 表，结构与 docx/xlsx 一致）；
* 图片 → ``image``（原图经 ``image_sink`` 落盘，对象键写进 ``Block.meta["image_key"]``）；
* 分组形状递归展开；演讲者备注 → ``para``，前缀 ``备注：``。

:param image_sink: ``(data, ext) -> key``。为 ``None`` 时**不抽取**内嵌图片
    （保持与原型一致的行为，parity 回归工具即走这条路径）。
"""

from __future__ import annotations

import logging
from collections.abc import Callable
from io import BytesIO

from app.core.parser.block import Block
from app.core.parser.markdown import to_md_table

logger = logging.getLogger(__name__)

ImageSink = Callable[[bytes, str], str]

#: python-pptx 未安装时的可执行提示
PPTX_MISSING_HINT = "未安装 python-pptx，无法解析 .pptx：pip install python-pptx"


def _para_texts(text_frame) -> list[str]:
    """文本框 → 段落文本（保留层级缩进；空段落丢弃）。"""
    out: list[str] = []
    for para in text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            # 有些文件把文字放在 para.text 而不在 run 里
            text = (para.text or "").strip()
        if not text:
            continue
        level = int(getattr(para, "level", 0) or 0)
        out.append(("  " * level) + text if level > 0 else text)
    return out


def _table_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return rows


def _walk_shapes(shapes, blocks: list[Block], image_sink: ImageSink | None, depth: int = 0) -> None:
    """按文档顺序遍历形状（分组递归；``depth`` 防御异常嵌套）。"""
    if depth > 6:
        return
    for shape in shapes:
        try:
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = _table_rows(shape.table)
                if rows:
                    blocks.append(Block(type="table", text=to_md_table(rows), rows=rows))
                continue

            shape_type = getattr(shape, "shape_type", None)
            type_name = getattr(shape_type, "name", str(shape_type or ""))

            if type_name == "GROUP":
                _walk_shapes(shape.shapes, blocks, image_sink, depth + 1)
                continue

            if type_name == "PICTURE" or getattr(shape, "image", None) is not None:
                meta: dict = {"image_name": getattr(shape, "name", "图片")}
                if image_sink is not None:
                    try:
                        blob = shape.image.blob
                        ext = f".{shape.image.ext}"
                        meta["image_key"] = image_sink(blob, ext)
                    except Exception as exc:  # noqa: BLE001 - 单张图失败不影响整篇
                        logger.warning("pptx 内嵌图片抽取失败：%s", exc)
                        meta["image_error"] = str(exc)
                blocks.append(Block(type="image", text=meta["image_name"], meta=meta))
                continue

            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts = _para_texts(shape.text_frame)
                if texts:
                    blocks.append(Block(type="para", text="\n".join(texts)))
        except Exception as exc:  # noqa: BLE001 - 单个形状异常不拖垮整篇
            logger.warning("pptx 形状解析跳过（%s）：%s", getattr(shape, "name", "?"), exc)


def pptx_to_blocks(data: bytes, *, image_sink: ImageSink | None = None) -> list[Block]:
    """``.pptx`` 字节 → 解析块。"""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - 依赖缺失时的业务提示
        raise RuntimeError(PPTX_MISSING_HINT) from exc

    prs = Presentation(BytesIO(data))
    blocks: list[Block] = []

    for index, slide in enumerate(prs.slides, start=1):
        blocks.append(Block(type="heading", level=1, text=f"幻灯片 {index}"))

        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:  # noqa: BLE001 - 部分文件标题占位符异常
            title = ""
        if title:
            blocks.append(Block(type="heading", level=2, text=title))

        _walk_shapes(slide.shapes, blocks, image_sink)

        notes = ""
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:  # noqa: BLE001 - 备注缺失是常态
            notes = ""
        if notes:
            blocks.append(Block(type="para", text=f"备注：{notes}"))

    return blocks


__all__ = ["PPTX_MISSING_HINT", "ImageSink", "pptx_to_blocks"]
